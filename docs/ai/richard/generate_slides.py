#!/usr/bin/env python3
"""
Generate richard_slides.pptx from richard.md and richard_commands.md plus figures in this folder.

Usage:
  python docs/ai/richard/generate_slides.py

Requires:
  python-pptx
"""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

BASE_DIR = Path(__file__).resolve().parent
MD_PATH = BASE_DIR / "richard.md"
CMD_PATH = BASE_DIR / "richard_commands.md"
IMG_DIR = BASE_DIR
OUT_PATH = BASE_DIR / "richard_slides.pptx"

IMAGE_MAP = {
    "Scope and data sources": ["richard_scope.png"],
    "Executive summary": ["richard_exec_summary.png"],
    "Dataset overview": ["richard_dataset_overview.png"],
    "Session structure": ["richard_session_structure.png"],
    "Intent taxonomy (fine-grained)": ["richard_intent_taxonomy.png"],
    "Category distribution (combined)": ["richard_category_distribution.png", "category_distribution.png"],
    "Semantic groupings": ["richard_semantic_groupings.png", "semantic_groupings_pie.png"],
    "Dominant session archetypes": ["richard_dominant_archetypes.png"],
    "Workflow flows (category sequences)": ["richard_workflow_flows.png", "workflow_2_step_flows_graph.png"],
    "Codebase focus (paths referenced in prompts)": ["richard_codebase_focus.png", "codebase_focus_treemap.png"],
    "Flow diagrams (ASCII)": ["richard_flow_diagrams.png"],
    "Prompt composition and specificity": ["richard_prompt_composition.png"],
    "Assistant response analysis": ["richard_response_analysis.png"],
    "Assistant failure modes and oversight signals": ["richard_failure_modes.png"],
    "Category transitions within sessions": ["richard_category_transitions.png"],
    "Temporal trends": ["richard_temporal_trends.png"],
    "Lexical themes (filtered prompts only)": ["richard_lexical_themes.png", "lexical_themes_unigrams_wordcloud.png"],
    "Tool and command mentions": ["richard_tool_mentions.png"],
    "Common prompt openers": ["richard_prompt_openers.png"],
    "Observed prompt archetypes": ["richard_archetypes.png"],
    "Investigation patterns (more in-depth)": ["richard_investigation_patterns.png"],
    "Prompt templates that match my usage": ["richard_prompt_templates.png"],
    "Claude Code usage (Metta workspaces)": ["richard_claude_usage.png"],
    "Limitations": ["richard_limitations.png"],
}


def parse_sections(text: str) -> list[dict]:
    lines = text.splitlines()
    sections = []
    current = None
    for line in lines:
        if line.startswith("## "):
            if current:
                sections.append(current)
            current = {"title": line[3:].strip(), "lines": []}
        elif current is not None:
            current["lines"].append(line)
    if current:
        sections.append(current)
    return sections


def pick_image(title: str) -> Path | None:
    for name in IMAGE_MAP.get(title, []):
        path = IMG_DIR / name
        if path.exists():
            return path
    return None


def extract_bullets(lines: list[str]) -> list[str]:
    bullets = []
    in_code = False
    for line in lines:
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            continue
        if line.startswith("- "):
            bullets.append(line[2:].strip())
    if bullets:
        return bullets[:6]
    for line in lines:
        s = line.strip()
        if not s or s.startswith("```") or s.startswith("|"):
            continue
        if s.startswith("#"):
            continue
        bullets.append(s)
        break
    if not bullets:
        bullets.append("Summary bullet not available.")
    return bullets


def main() -> None:
    sections = parse_sections(MD_PATH.read_text())
    ppt = Presentation()

    # Title slide
    slide = ppt.slides.add_slide(ppt.slide_layouts[0])
    slide.shapes.title.text = "Richard Higgins – Codex Usage Analysis"
    slide.placeholders[1].text = "Updated from docs/ai/richard/richard.md and richard_commands.md"

    # Layout constants
    margin = Inches(0.5)
    text_width = Inches(5.6)
    img_left = Inches(6.1)
    img_top = Inches(1.4)
    img_width = Inches(6.7)

    for section in sections:
        title = section["title"]
        bullets = extract_bullets(section["lines"])
        img_path = pick_image(title)

        slide = ppt.slides.add_slide(ppt.slide_layouts[1])  # Title + Content
        slide.shapes.title.text = title
        body = slide.shapes.placeholders[1]
        body.left = margin
        body.top = Inches(1.4)
        body.width = text_width
        body.height = Inches(5.5)

        tf = body.text_frame
        tf.clear()
        for i, bullet in enumerate(bullets):
            p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
            p.text = bullet
            p.level = 0

        if img_path:
            slide.shapes.add_picture(str(img_path), img_left, img_top, width=img_width)

        if title == "Lexical themes (filtered prompts only)":
            uni = IMG_DIR / "lexical_themes_unigrams_wordcloud.png"
            bi = IMG_DIR / "lexical_themes_bigrams_wordcloud.png"
            if uni.exists() or bi.exists():
                slide_wc = ppt.slides.add_slide(ppt.slide_layouts[5])
                slide_wc.shapes.title.text = "Lexical Themes – Wordclouds"
                if uni.exists():
                    slide_wc.shapes.add_picture(str(uni), Inches(0.6), Inches(1.4), width=Inches(6.5))
                if bi.exists():
                    slide_wc.shapes.add_picture(str(bi), Inches(7.0), Inches(1.4), width=Inches(6.5))

        if title == "Workflow flows (category sequences)":
            w2 = IMG_DIR / "workflow_2_step_flows_graph.png"
            w3 = IMG_DIR / "workflow_3_step_flows_barchart.png"
            if w2.exists():
                slide_w2 = ppt.slides.add_slide(ppt.slide_layouts[5])
                slide_w2.shapes.title.text = "Workflow Flows – 2-Step"
                slide_w2.shapes.add_picture(str(w2), margin, Inches(1.4), width=Inches(12.5))
            if w3.exists():
                slide_w3 = ppt.slides.add_slide(ppt.slide_layouts[5])
                slide_w3.shapes.title.text = "Workflow Flows – 3-Step"
                slide_w3.shapes.add_picture(str(w3), margin, Inches(1.4), width=Inches(12.5))

    # Commands section
    commands = []
    current = None
    for line in CMD_PATH.read_text().splitlines():
        line = line.strip()
        if line.startswith("### /"):
            if current:
                commands.append(current)
            current = {"name": line.replace("### ", ""), "purpose": ""}
        elif line.startswith("Purpose:") and current is not None:
            current["purpose"] = line.replace("Purpose:", "").strip()
    if current:
        commands.append(current)

    slide = ppt.slides.add_slide(ppt.slide_layouts[0])
    slide.shapes.title.text = "Workspace-Agnostic Slash Commands"
    slide.placeholders[1].text = "Summary from docs/ai/richard/richard_commands.md"

    chunk = 6
    for i in range(0, len(commands), chunk):
        group = commands[i : i + chunk]
        slide = ppt.slides.add_slide(ppt.slide_layouts[1])
        slide.shapes.title.text = "Slash Commands"
        body = slide.shapes.placeholders[1].text_frame
        body.clear()
        for j, cmd in enumerate(group):
            text = f"{cmd['name']} — {cmd['purpose']}" if cmd["purpose"] else cmd["name"]
            p = body.add_paragraph() if j > 0 else body.paragraphs[0]
            p.text = text
            p.level = 0

    ppt.save(OUT_PATH)


if __name__ == "__main__":
    main()
